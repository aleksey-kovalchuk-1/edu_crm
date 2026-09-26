import io
import json
from dataclasses import replace

import openpyxl
import pytest
from cryptography.fernet import Fernet
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from app import models
from app import customer_imports
from app import customer_import_routes
from app.customer_imports import read_rows
from app.security import TokenCipher
from helpers import database, login


@pytest.fixture
def head(app, keycloak):
    app.state.learner_cipher = TokenCipher(Fernet.generate_key())
    with TestClient(app) as client:
        login(client, keycloak, roles=('crm-supervisor',), subject='kc-import-head')
        yield client


def upload(client, kind, phase, content, filename):
    return client.post(f'/api/v1/customer-imports/{kind}/{phase}', files={'file': (filename, content)})


def workbook(headers, rows):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def applications():
    return [
        {'Номер заявки': f'A-{number}', 'Курс': 'Python', 'Фамилия': 'Тестов',
         'Имя': 'Иван', 'Отчество': 'Иванович', 'Телефон': '+7 (900) 000-00-01',
         'Email': 'learner@example.test', 'Номер потока': f'P-{number}'}
        for number in range(1, 6)
    ] + [None]


def test_application_json_previews_five_rows_and_skips_null_idempotently(head, database_url):
    content = json.dumps(applications(), ensure_ascii=False).encode()
    preview = upload(head, 'applications', 'preview', content, 'Данные оплат.json')
    assert preview.status_code == 200, preview.text
    assert preview.json()['summary']['valid'] == 5
    assert preview.json()['summary']['skipped'] == 1
    assert any('null' in warning.lower() for row in preview.json()['rows'] for warning in row['warnings'])
    with database(database_url) as db:
        assert db.scalar(select(func.count()).select_from(models.CourseApplication)) == 0
    applied = upload(head, 'applications', 'apply', content, 'Данные оплат.json')
    assert applied.status_code == 200, applied.text
    assert applied.json()['summary']['created'] == 5
    again = upload(head, 'applications', 'apply', content, 'Данные оплат.json')
    assert again.status_code == 200
    assert again.json()['summary']['created'] == 0
    with database(database_url) as db:
        apps = db.scalars(select(models.CourseApplication)).all()
        assert len(apps) == 5
        assert len({item.learner_id for item in apps}) == 1
        assert {item.payment_status for item in apps} == {'unconfirmed_by_data'}


def test_ambiguous_phone_match_is_reported_without_merging(head, database_url):
    first = head.post('/api/v1/learners', json={'last_name': 'Один', 'first_name': 'Тест', 'phone': '79000000001'}).json()
    second = head.post('/api/v1/learners', json={'last_name': 'Два', 'first_name': 'Тест', 'phone': '+7 900 000 00 01'}).json()
    content = json.dumps([applications()[0]], ensure_ascii=False).encode()
    preview = upload(head, 'applications', 'preview', content, 'заявки.json')
    assert preview.status_code == 200
    assert preview.json()['summary']['invalid'] == 1
    assert 'неоднознач' in preview.json()['rows'][0]['errors'][0].lower()
    assert set(preview.json()['rows'][0]['candidate_ids']) == {first['id'], second['id']}
    assert preview.json()['rows'][0]['signals'][0]['rule_code'] == 'learner_match_conflict'
    applied = upload(head, 'applications', 'apply', content, 'заявки.json')
    assert applied.json()['summary']['created'] == 0
    with database(database_url) as db:
        assert db.scalar(select(func.count()).select_from(models.CourseApplication)) == 0
    resolved = head.post('/api/v1/customer-imports/applications/apply',
                         files={'file': ('заявки.json', content)},
                         data={'resolved_learner_ids': json.dumps({'1': first['id']})})
    assert resolved.status_code == 200, resolved.text
    assert resolved.json()['summary']['created'] == 1
    assert head.get('/api/v1/course-applications').json()[0]['learner_id'] == first['id']


def test_shared_family_phone_does_not_merge_distinct_learners(head):
    first = head.post('/api/v1/learners', json={'last_name': 'Один', 'first_name': 'Тест',
                                               'phone': '79000000001', 'email': 'one@example.test'}).json()
    content = workbook(['Фамилия', 'Имя', 'Телефон', 'Email'],
                       [['Другой', 'Тест', '79000000001', 'two@example.test']])
    result = upload(head, 'learners', 'apply', content, 'synthetic.xlsx').json()
    assert result['summary']['created'] == 1
    assert result['rows'][0]['signals'][0]['rule_code'] == 'shared_contact'
    assert result['rows'][0]['warnings']
    assert len(head.get('/api/v1/learners').json()) == 2
    assert result['record_links'][0]['entity_id'] != first['id']


def test_application_with_shared_phone_keeps_distinct_learner_names(head):
    first = head.post('/api/v1/learners', json={'last_name': 'Один', 'first_name': 'Тест',
                                               'phone': '79000000001'}).json()
    item = {**applications()[0], 'Фамилия': 'Другой', 'Имя': 'Слушатель'}
    content = json.dumps([item], ensure_ascii=False).encode()
    preview = upload(head, 'applications', 'preview', content, 'synthetic.json').json()
    assert preview['rows'][0]['signals'][0]['rule_code'] == 'shared_contact'
    applied = upload(head, 'applications', 'apply', content, 'synthetic.json').json()
    assert applied['summary']['created'] == 1
    assert head.get('/api/v1/course-applications').json()[0]['learner_name'] == 'Другой Слушатель Иванович'
    assert len(head.get('/api/v1/learners').json()) == 2
    assert head.get('/api/v1/course-applications').json()[0]['learner_id'] != first['id']


def test_existing_application_rejects_changed_name_even_with_same_contact(head):
    item = applications()[0]
    content = json.dumps([item], ensure_ascii=False).encode()
    assert upload(head, 'applications', 'apply', content, 'synthetic.json').status_code == 200
    changed = {**item, 'Фамилия': 'Подмененный'}
    result = upload(head, 'applications', 'apply', json.dumps([changed], ensure_ascii=False).encode(),
                    'synthetic.json').json()
    assert result['summary']['invalid'] == 1
    assert result['rows'][0]['signals'][0]['rule_code'] == 'application_number_conflict'
    assert head.get('/api/v1/course-applications').json()[0]['learner_name'] == 'Тестов Иван Иванович'


def test_import_velocity_warns_without_blocking_rows(head):
    head.app.state.settings = replace(head.app.state.settings, fraud_batch_row_limit=1, fraud_hourly_import_limit=1)
    content = json.dumps([applications()[0], None], ensure_ascii=False).encode()
    preview = upload(head, 'applications', 'preview', content, 'synthetic.json').json()
    assert preview['batch_signals'][0]['rule_code'] == 'import_velocity'
    first = upload(head, 'applications', 'apply', content, 'synthetic.json').json()
    assert first['summary']['created'] == 1
    second = upload(head, 'applications', 'apply', content, 'synthetic.json').json()
    assert second['batch_signals'][0]['rule_code'] == 'import_velocity'


def test_duplicate_external_number_is_reported_and_second_row_is_skipped(head):
    repeated = [applications()[0], {**applications()[0], 'Курс': 'Другой курс'}]
    content = json.dumps(repeated, ensure_ascii=False).encode()
    preview = upload(head, 'applications', 'preview', content, 'заявки.json').json()
    assert preview['summary']['invalid'] == 1
    assert 'повторяется' in preview['rows'][1]['errors'][0]
    assert preview['rows'][1]['signals'][0]['rule_code'] == 'batch_repetition'
    assert upload(head, 'applications', 'apply', content, 'заявки.json').json()['summary']['created'] == 1
    assert head.get('/api/v1/course-applications').json()[0]['course'] == 'Python'


def test_reimported_application_number_cannot_overwrite_conflicting_course(head):
    first = applications()[0]
    assert upload(head, 'applications', 'apply', json.dumps([first], ensure_ascii=False).encode(), 'demo.json').status_code == 200
    changed = {**first, 'Курс': 'Другой курс', 'Номер потока': 'P-2'}
    preview = upload(head, 'applications', 'preview', json.dumps([changed], ensure_ascii=False).encode(), 'demo.json').json()
    assert preview['summary']['invalid'] == 1
    assert preview['rows'][0]['signals'][0]['rule_code'] == 'application_number_conflict'
    applied = upload(head, 'applications', 'apply', json.dumps([changed], ensure_ascii=False).encode(), 'demo.json').json()
    assert applied['summary']['updated'] == 0
    card = head.get('/api/v1/course-applications').json()[0]
    assert card['course'] == 'Python'
    assert card['stream_number'] == 'P-1'


def test_reimported_application_number_cannot_change_learner(head):
    first = applications()[0]
    upload(head, 'applications', 'apply', json.dumps([first], ensure_ascii=False).encode(), 'demo.json')
    changed = {**first, 'Фамилия': 'Другой', 'Имя': 'Человек',
               'Телефон': '79000000002', 'Email': 'other@example.test'}
    response = upload(head, 'applications', 'apply', json.dumps([changed], ensure_ascii=False).encode(), 'demo.json')
    assert response.json()['summary']['invalid'] == 1
    assert response.json()['rows'][0]['signals'][0]['rule_code'] == 'application_number_conflict'
    assert head.get('/api/v1/course-applications').json()[0]['learner_name'] == 'Тестов Иван Иванович'


def test_vendor_table_splits_products_and_reuses_company_contact(head):
    content = workbook(
        ['Компания', 'Продукт', 'ФИО', 'Телефон', 'Почта', 'Способ связи'],
        [['ООО «ТДата»', 'RT.DataLake, RT.Warehouse', 'Тестовый Контакт', '00123',
          'contact@example.test', 'Почта; Чат в ТГ']],
    )
    preview = upload(head, 'vendors', 'preview', content, 'поставщики.xlsx')
    assert preview.status_code == 200, preview.text
    assert preview.json()['summary']['valid'] == 1
    applied = upload(head, 'vendors', 'apply', content, 'поставщики.xlsx')
    assert applied.status_code == 200, applied.text
    companies = head.get('/api/v1/vendor-companies').json()
    assert [company['name'] for company in companies] == ['ООО «ТДата»']
    contacts = head.get('/api/v1/vendor-contacts').json()
    assert len(contacts) == 1
    assert len(contacts[0]['product_ids']) == 2
    assert contacts[0]['preferred_channels'] == ['Почта', 'Чат в ТГ']
    assert upload(head, 'vendors', 'apply', content, 'поставщики.xlsx').json()['summary']['created'] == 0


def test_learner_reimport_does_not_erase_filled_fields_with_blank_cells(head):
    headers = ['Фамилия', 'Имя', 'Телефон', 'Email', 'Образование', 'Серия паспорта']
    initial = workbook(headers, [['Тестов', 'Иван', '00123', 'learner@example.test', 'Высшее', '0011']])
    assert upload(head, 'learners', 'apply', initial, 'слушатели.xlsx').status_code == 200
    later = workbook(headers, [['Тестов', 'Иван', '00123', 'learner@example.test', '', '']])
    assert upload(head, 'learners', 'apply', later, 'слушатели.xlsx').status_code == 200
    learner_id = head.get('/api/v1/learners').json()[0]['id']
    detail = head.get(f'/api/v1/learners/{learner_id}').json()
    assert detail['education'] == 'Высшее'
    assert detail['passport_series'] == '0011'


def test_numeric_phone_cell_is_rejected_before_leading_zero_can_be_lost(head):
    content = workbook(['Фамилия', 'Имя', 'Телефон'], [['Тестов', 'Иван', 12345]])
    preview = upload(head, 'learners', 'preview', content, 'слушатели.xlsx').json()
    assert preview['summary']['invalid'] == 1
    assert 'текстом' in preview['rows'][0]['errors'][0]


@pytest.mark.parametrize('number', [79000000001, 89000000001])
def test_numeric_eleven_digit_phone_is_accepted_with_row_warning(head, number):
    content = workbook(['Фамилия', 'Имя', 'Номер телефона'], [['Тестов', 'Иван', number]])
    preview = upload(head, 'learners', 'preview', content, 'synthetic.xlsx').json()
    assert preview['summary']['valid'] == 1
    assert len(preview['rows'][0]['warnings']) == 1
    assert str(number) not in json.dumps(preview, ensure_ascii=False)


@pytest.mark.parametrize('number', [9000000001, 90000000001, 79000000001.5])
def test_uncertain_numeric_phone_is_rejected(number):
    with pytest.raises(customer_imports.ImportRowError):
        customer_imports.normalize_import_phone(number)


def test_exact_customer_learner_template_maps_all_named_columns():
    headers = [
        'Фамилия', 'Имя', 'Отчествопри наличии)', 'Номер телефона', 'Email', 'СНИЛС',
        'Серия паспорта', 'Номер паспорта', 'Кем выдан паспорт', 'Дата выдачи паспорта',
        'Код подразделения', 'Пол', 'Дата рождения', 'Регион регистрации',
        'Населенный пункт регистрации', 'Улица регистрации', 'Дом регистрации',
        'Квартира регистрации', 'Индекс регистрации', 'Имядательный падеж)',
        'Фамилиядательный падеж)', 'Отчестводательный падеж)', 'Образование',
        'Профессия по диплому', 'Учебное заведение по диплому',
        'Фамилия, указанная в дипломе', 'Номер диплома', 'Серия диплома',
        'Регистрационный номер диплома', 'Дата выдачи диплома', None,
    ]
    values = [''] * len(headers)
    values[0], values[1], values[3] = 'Тестов', 'Иван', 79000000001
    book = openpyxl.Workbook()
    book.active.append(headers)
    book.active.append(values)
    book.create_sheet('Вспомогательный').append(['Не импортировать'])
    content = io.BytesIO()
    book.save(content)
    parsed = customer_imports.read_customer_file('learners', 'synthetic.xlsx', content.getvalue())
    assert parsed.template_version == 'customer-learners-v1'
    assert len(parsed.mapping) == 30
    assert parsed.unmapped_headers == []
    assert len(parsed.rows) == 1
    assert len(parsed.rows[0][1]) == 30
    assert parsed.rows[0][1]['phone'] == 79000000001


def test_customer_file_describes_vendor_mapping_without_source_values():
    content = workbook(['Компания', 'Продукт', 'ФИО', 'Телефон', 'Почта', 'Способ связи'],
                       [['Демо компания', 'Демо продукт', 'Тестовый Контакт', '79000000001',
                         'demo@example.test', 'Почта']])
    parsed = customer_imports.read_customer_file('vendors', 'synthetic.xlsx', content)
    assert parsed.template_version == 'customer-vendors-v1'
    assert parsed.mapping['company'] == 'Компания'
    assert parsed.unmapped_headers == []
    assert len(parsed.rows) == 1


def test_customer_header_only_preview_and_apply_rejection(head):
    content = workbook(['Фамилия', 'Имя'], [])
    preview = upload(head, 'learners', 'preview', content, 'blank.xlsx')
    assert preview.status_code == 200, preview.text
    assert preview.json()['summary']['rows'] == 0
    assert preview.json()['mapping'] == {'last_name': 'Фамилия', 'first_name': 'Имя'}
    assert upload(head, 'learners', 'apply', content, 'blank.xlsx').status_code == 422


def test_unknown_column_and_manual_mapping(head):
    content = workbook(['Фамилия', 'Имя', 'Контактный телефон'], [['Тестов', 'Иван', '79000000001']])
    preview = upload(head, 'learners', 'preview', content, 'synthetic.xlsx').json()
    assert preview['unmapped_headers'] == ['Контактный телефон']
    assert preview['summary']['invalid'] == 1
    mapping = {'last_name': 'Фамилия', 'first_name': 'Имя', 'phone': 'Контактный телефон'}
    response = head.post('/api/v1/customer-imports/learners/preview',
                         files={'file': ('synthetic.xlsx', content)}, data={'mapping': json.dumps(mapping)})
    assert response.status_code == 200, response.text
    assert response.json()['summary']['valid'] == 1
    assert response.json()['unmapped_headers'] == []
    bad = head.post('/api/v1/customer-imports/learners/preview',
                    files={'file': ('synthetic.xlsx', content)}, data={'mapping': json.dumps({'phone': 'Чужой столбец'})})
    assert bad.status_code == 422


def test_customer_import_is_supervisor_only(app, keycloak):
    content = workbook(['Фамилия', 'Имя'], [])
    with TestClient(app) as client:
        login(client, keycloak, roles=('crm-user',), subject='kc-import-regular')
        assert upload(client, 'learners', 'preview', content, 'blank.xlsx').status_code == 403
        assert upload(client, 'learners', 'apply', content, 'blank.xlsx').status_code == 403


def test_duplicate_target_previews_but_apply_needs_manual_mapping(head):
    content = workbook(['Фамилия', 'Имя', 'Телефон', 'Номер телефона'],
                       [['Тестов', 'Иван', '79000000001', '79000000001']])
    preview = upload(head, 'learners', 'preview', content, 'synthetic.xlsx')
    assert preview.status_code == 200, preview.text
    assert preview.json()['mapping_conflicts'] == {'phone': ['Телефон', 'Номер телефона']}
    assert upload(head, 'learners', 'apply', content, 'synthetic.xlsx').status_code == 422
    parsed = customer_imports.read_customer_file('learners', 'synthetic.xlsx', content,
                                                 selected_mapping={'last_name': 'Фамилия', 'first_name': 'Имя',
                                                                   'phone': 'Номер телефона'})
    assert parsed.unmapped_headers == ['Телефон']
    assert parsed.mapping_conflicts == {}
    mapped = head.post('/api/v1/customer-imports/learners/apply',
                       files={'file': ('synthetic.xlsx', content)},
                       data={'mapping': json.dumps(parsed.mapping)})
    assert mapped.status_code == 200, mapped.text
    assert mapped.json()['summary']['created'] == 1


def test_apply_records_safe_batch_history_and_card_link(head, database_url):
    content = workbook(['Фамилия', 'Имя', 'Телефон'], [['Тестов', 'Иван', '79000000001']])
    response = upload(head, 'learners', 'apply', content, 'private-name.xlsx')
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['batch_id'] > 0
    assert result['record_links'][0]['entity_type'] == 'learner'
    learner_id = result['record_links'][0]['entity_id']
    assert head.get(f'/api/v1/learners/{learner_id}').status_code == 200
    history = head.get('/api/v1/customer-imports/history')
    assert history.status_code == 200
    assert history.json()[0]['id'] == result['batch_id']
    detail = head.get(f"/api/v1/customer-imports/history/{result['batch_id']}")
    assert detail.status_code == 200
    assert detail.json()['record_links'][0]['entity_id'] == learner_id
    with database(database_url) as db:
        audit = db.scalars(select(models.AuditEvent)).all()
        assert '79000000001' not in json.dumps([event.payload for event in audit])
    assert 'private-name.xlsx' not in json.dumps(history.json())
    assert '79000000001' not in json.dumps(detail.json())


def test_apply_failure_rolls_back_cards_and_batch(head, database_url, monkeypatch):
    def fail_after_rows(*args, **kwargs):
        raise RuntimeError('synthetic failure')
    monkeypatch.setattr(customer_import_routes, 'record_event', fail_after_rows)
    content = workbook(['Фамилия', 'Имя', 'Телефон'], [['Тестов', 'Иван', '79000000001']])
    with pytest.raises(RuntimeError, match='synthetic failure'):
        upload(head, 'learners', 'apply', content, 'synthetic.xlsx')
    with database(database_url) as db:
        assert db.scalar(select(func.count()).select_from(models.Learner)) == 0
        assert db.scalar(select(func.count()).select_from(models.CustomerImportBatch)) == 0
        assert db.scalar(select(func.count()).select_from(models.CustomerImportRowLink)) == 0


def test_learner_import_accepts_russian_date_format(head):
    content = workbook(['Фамилия', 'Имя', 'Телефон', 'Дата рождения'],
                       [['Тестов', 'Иван', '00123', '01.02.2000']])
    applied = upload(head, 'learners', 'apply', content, 'слушатели.xlsx')
    assert applied.status_code == 200
    assert applied.json()['summary']['created'] == 1
    learner_id = head.get('/api/v1/learners').json()[0]['id']
    assert head.get(f'/api/v1/learners/{learner_id}').json()['birth_date'] == '2000-02-01'
