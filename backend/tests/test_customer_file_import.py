"""«Загрузка справочников» accepts the customer's applications JSON and the customer's workbook (owner 29.09.2026, D-247).

The JSON keeps only the application number, course and stream: names, phones and e-mails are dropped when the file
is read and never reach the database. The workbook is read by sheet name; sheets without a matching CRM entity are
reported, never imported elsewhere. Universities are not created from the workbook.
"""
import io
import json

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, inspect, select, text

from app.models import AuditEvent, CourseApplication, ITDirection, ITProduct, University, VendorCompany
from helpers import database, login

XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
PII = ('Черепанова-тест', '7 (999) 000-00-01', 'pii.marker@test.local', 'Отчество-тест')


def applications_json(entries=None):
    if entries is None:
        entries = [
            None,
            {'Номер заявки': 'ORD-TEST-0001', 'Курс': 'Анализ данных без программирования', 'Фамилия': PII[0],
             'Имя': 'Светлана', 'Отчество': PII[3], 'Телефон': PII[1], 'Email': PII[2], 'Номер потока': 1},
            {'Номер заявки': 'ORD-TEST-0002', 'Курс': 'Инженер-тестировщик', 'Фамилия': 'Демо', 'Имя': 'Максим',
             'Отчество': '', 'Телефон': '7 (997) 000-00-02', 'Email': 'demo2@test.local', 'Номер потока': 2},
        ]
    return json.dumps(entries, ensure_ascii=False).encode()


def _sheet(book, title, headers, rows, title_rows=('Пилотная книга', 'Описание листа')):
    sheet = book.create_sheet(title)
    for line in title_rows:
        sheet.append([line])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)


def customer_workbook(universities=None, directions=None, products=None):
    book = openpyxl.Workbook()
    book.active.title = 'Сводка'
    book['Сводка'].append(['Вузов', 'Программ в пилоте'])
    book['Сводка'].append([20, 80])
    _sheet(book, 'Вузы', ['university_id', 'Полное наименование', 'Краткое наименование', 'Регион', 'Город', 'Официальный сайт'],
           universities if universities is not None else [
               ['UNI-001', 'Московский физико-технический институт', 'МФТИ', 'Московская область', 'Долгопрудный', 'https://mipt.ru'],
               ['UNI-002', 'Неизвестный университет', 'НУ', 'Регион', 'Город', 'https://example.test'],
           ])
    _sheet(book, 'Программы', ['program_id', 'university_id', 'Название программы'], [['PRG-001', 'UNI-001', 'Прикладная математика']])
    _sheet(book, 'Продукты РТК', ['product_id', 'Наименование', 'Тип', 'Описание'],
           products if products is not None else [['RTK-LMS-001', 'СДО ИТ Школы РТК', 'Внешняя LMS-платформа', 'Существующая LMS']])
    _sheet(book, 'Направления', ['direction_id', 'Направление', 'Содержание', 'Покрытие РТК'],
           directions if directions is not None else [['DIR-001', 'Анализ данных', 'Методы анализа данных', 'Да']])
    _sheet(book, 'Ответственные', ['assignment_id', 'university_id', 'ФИО ответственного', 'Корпоративный email'], [['A-1', 'UNI-001', None, None]])
    _sheet(book, 'Договоры', ['contract_id', 'university_id', 'Номер', 'Статус'], [[None, 'UNI-001', None, 'Не проверено']])
    _sheet(book, 'Продукты договора', ['contract_product_id', 'contract_id', 'product_id'], [[None, None, None]])
    _sheet(book, 'Допсоглашения', ['addendum_id', 'contract_id', 'Номер'], [[None, None, None]])
    _sheet(book, 'Источники', ['source_id', 'Владелец', 'URL'], [['SRC-1', 'Вуз', 'https://example.test']])
    _sheet(book, 'Справочники', ['Справочник', 'Значение', 'Описание'], [['match_status', 'Да', 'Прямое соответствие']])
    _sheet(book, 'Инструкция', ['Раздел', 'Правило'], [['1. Назначение', 'Книга для импорта']])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def head(app, keycloak):
    with TestClient(app) as client:
        login(client, keycloak, roles=('crm-supervisor',), subject='kc-head', name='Павел Демо', email='pavel.demo@demo.local')
        yield client


@pytest.fixture
def mipt(database_url):
    with database(database_url) as db:
        university = University(name='МФТИ', city='', contact='')
        db.add(university)
        db.commit()
        return university.id


def upload(client, content, filename):
    content_type = 'application/json' if filename.endswith('.json') else XLSX_TYPE
    return client.post('/api/v1/imports', files={'file': (filename, content, content_type)})


def check(client, import_id):
    return client.post(f'/api/v1/imports/{import_id}/check', json={'mapping': {}})


def apply(client, import_id):
    return client.post(f'/api/v1/imports/{import_id}/apply', json={'mapping': {}})


def everything_as_text(database_url):
    engine = create_engine(database_url)
    try:
        with engine.connect() as connection:
            names = inspect(connection).get_table_names()
            return ' '.join(str(connection.execute(text(f'select coalesce(json_agg(t)::text, \'\') from "{name}" t')).scalar_one())
                            for name in names)
    finally:
        engine.dispose()


# ---- applications JSON -------------------------------------------------------------------------------------------

def test_the_json_upload_keeps_only_number_course_and_stream(head, database_url):
    response = upload(head, applications_json(), 'Данные оплат.json')
    assert response.status_code == 201, response.text
    body = response.json()
    assert body['kind'] == 'applications'
    assert body['headers'] == ['Номер заявки', 'Курс', 'Номер потока']
    assert body['row_count'] == 3
    assert body['preview'][1]['cells'] == ['ORD-TEST-0001', 'Анализ данных без программирования', '1']
    assert not any(marker in response.text for marker in PII)
    stored = everything_as_text(database_url)
    assert not any(marker in stored for marker in PII)


def test_check_previews_without_writing_and_reports_null_entries(head, database_url):
    import_id = upload(head, applications_json(), 'Данные оплат.json').json()['id']
    report = check(head, import_id).json()
    assert report['summary']['rows'] == 3
    assert report['summary']['valid'] == 2
    assert report['summary']['skipped'] == 1
    assert report['summary']['created'] == {'course_applications': 2}
    skipped = [row for row in report['rows'] if row['status'] == 'skipped']
    assert skipped[0]['row_number'] == 1 and 'null' in skipped[0]['warnings'][0]
    with database(database_url) as db:
        assert db.scalars(select(CourseApplication)).all() == []


def test_apply_creates_applications_without_a_learner_and_without_payment(head, database_url):
    import_id = upload(head, applications_json(), 'Данные оплат.json').json()['id']
    response = apply(head, import_id)
    assert response.status_code == 200, response.text
    with database(database_url) as db:
        rows = {row.external_number: row for row in db.scalars(select(CourseApplication))}
        assert set(rows) == {'ORD-TEST-0001', 'ORD-TEST-0002'}
        assert rows['ORD-TEST-0001'].learner_id is None
        assert rows['ORD-TEST-0001'].stream_number == '1'
        assert {row.payment_status for row in rows.values()} == {'unconfirmed_by_data'}
        event = db.scalar(select(AuditEvent).where(AuditEvent.action == 'import.apply'))
        assert event.payload['summary']['created'] == {'course_applications': 2}
    assert not any(marker in everything_as_text(database_url) for marker in PII)


def test_a_reupload_updates_by_application_number_without_duplicates(head, database_url):
    apply(head, upload(head, applications_json(), 'Данные оплат.json').json()['id'])
    changed = applications_json([
        {'Номер заявки': 'ORD-TEST-0001', 'Курс': 'Анализ данных (новая редакция)', 'Номер потока': 5},
        {'Номер заявки': 'ORD-TEST-0002', 'Курс': 'Инженер-тестировщик', 'Номер потока': 2},
    ])
    import_id = upload(head, changed, 'Данные оплат.json').json()['id']
    report = check(head, import_id).json()
    assert report['summary']['created'] == {'course_applications': 0}
    assert report['summary']['updated'] == {'course_applications': 1}
    assert [row['action'] for row in report['rows']] == ['update', 'unchanged']
    apply(head, import_id)
    with database(database_url) as db:
        rows = db.scalars(select(CourseApplication)).all()
        assert len(rows) == 2
        first = next(row for row in rows if row.external_number == 'ORD-TEST-0001')
        assert (first.course, first.stream_number) == ('Анализ данных (новая редакция)', '5')


@pytest.mark.parametrize('entry,reason', [
    ({'Курс': 'Курс', 'Номер потока': 1}, 'Номер заявки'),
    ({'Номер заявки': 'ORD-X', 'Номер потока': 1}, 'Курс'),
    ({'Номер заявки': 'ORD-X', 'Курс': 'Курс'}, 'Номер потока'),
    ({'Номер заявки': 'ORD-X', 'Курс': 'Курс', 'Номер потока': True}, 'Номер потока'),
    ({'Номер заявки': ['ORD-X'], 'Курс': 'Курс', 'Номер потока': 1}, 'Номер заявки'),
    ('строка вместо записи', 'не является объектом'),
])
def test_an_invalid_entry_is_skipped_with_a_reason(head, database_url, entry, reason):
    good = {'Номер заявки': 'ORD-OK', 'Курс': 'Курс', 'Номер потока': 1}
    import_id = upload(head, applications_json([entry, good]), 'заявки.json').json()['id']
    report = apply(head, import_id).json()
    bad = report['rows'][0]
    assert bad['status'] == 'error'
    assert reason in bad['errors'][0]
    assert report['summary']['created'] == {'course_applications': 1}


def test_a_number_repeated_in_the_file_uses_the_last_entry(head, database_url):
    entries = [{'Номер заявки': 'ORD-R', 'Курс': 'Первый', 'Номер потока': 1},
               {'Номер заявки': 'ORD-R', 'Курс': 'Второй', 'Номер потока': 2}]
    report = apply(head, upload(head, applications_json(entries), 'заявки.json').json()['id']).json()
    assert [row['status'] for row in report['rows']] == ['skipped', 'ok']
    with database(database_url) as db:
        assert [row.course for row in db.scalars(select(CourseApplication))] == ['Второй']


@pytest.mark.parametrize('content,message', [
    (b'{"a": 1}', 'массив'),
    (b'[1, 2', 'JSON'),
    (json.dumps([{'Имя': 'x'}]).encode(), 'Номер заявки'),
])
def test_a_json_file_without_applications_is_refused(head, content, message):
    response = upload(head, content, 'заявки.json')
    assert response.status_code == 422
    assert message in response.json()['message']


# ---- customer workbook -----------------------------------------------------------------------------------------

def test_the_workbook_is_read_by_sheet_name_and_lists_every_sheet(head, mipt):
    response = upload(head, customer_workbook(), 'RTK_IT_School_CRM_20.xlsx')
    assert response.status_code == 201, response.text
    body = response.json()
    assert body['kind'] == 'workbook'
    sheets = {sheet['name']: sheet for sheet in body['sheets']}
    assert list(sheets) == ['Сводка', 'Вузы', 'Программы', 'Продукты РТК', 'Направления', 'Ответственные', 'Договоры',
                            'Продукты договора', 'Допсоглашения', 'Источники', 'Справочники', 'Инструкция']
    assert {name for name, sheet in sheets.items() if sheet['status'] == 'supported'} == {'Вузы', 'Направления', 'Продукты РТК'}
    assert sheets['Программы']['status'] == 'not_supported' and 'нет сущности' in sheets['Программы']['reason']
    assert sheets['Договоры']['status'] == 'not_supported' and 'не заполнен' in sheets['Договоры']['reason']
    assert sheets['Сводка']['status'] == 'service'
    assert sheets['Вузы']['rows'] == 2
    # Title rows above the header are not counted as data rows.
    assert sheets['Программы']['rows'] == 1
    assert sheets['Источники']['rows'] == 1


def test_check_lists_unmatched_universities_and_apply_skips_them(head, database_url, mipt):
    import_id = upload(head, customer_workbook(), 'RTK_IT_School_CRM_20.xlsx').json()['id']
    report = check(head, import_id).json()
    assert report['unmatched_universities'] == [{'external_id': 'UNI-002', 'name': 'Неизвестный университет', 'short_name': 'НУ'}]
    assert report['summary']['created'] == {'universities': 0, 'it_directions': 1, 'it_products': 1}
    assert report['summary']['updated']['universities'] == 1
    applied = apply(head, import_id).json()
    skipped = [row for row in applied['rows'] if row['status'] == 'skipped']
    assert [(row['sheet'], row['key']) for row in skipped] == [('Вузы', 'UNI-002')]
    assert 'не найден в каталоге' in skipped[0]['warnings'][0]
    with database(database_url) as db:
        assert db.scalar(select(University).where(University.name == 'Неизвестный университет')) is None
        university = db.get(University, mipt)
        assert university.name == 'МФТИ'  # the CRM name is kept
        assert (university.external_id, university.city, university.region, university.website) == (
            'UNI-001', 'Долгопрудный', 'Московская область', 'https://mipt.ru')
        product = db.scalar(select(ITProduct).where(ITProduct.external_id == 'RTK-LMS-001'))
        assert (product.vendor, product.name, product.description) == ('ПАО «Ростелеком»', 'СДО ИТ Школы РТК', 'Существующая LMS')
        assert product.company.name == 'ПАО «Ростелеком»'
        direction = db.scalar(select(ITDirection).where(ITDirection.external_id == 'DIR-001'))
        assert (direction.name, direction.description) == ('Анализ данных', 'Методы анализа данных')


def test_a_workbook_reupload_updates_by_stable_ids(head, database_url, mipt):
    apply(head, upload(head, customer_workbook(), 'RTK.xlsx').json()['id'])
    again = customer_workbook(
        universities=[['UNI-001', 'Другое полное имя', 'МФТИ (новое)', 'Московская область', 'Москва', 'https://mipt.ru']],
        directions=[['DIR-001', 'Анализ данных и ИИ', 'Обновлено', 'Да']],
        products=[['RTK-LMS-001', 'СДО ИТ Школы', 'Внешняя LMS-платформа', 'Обновлено']],
    )
    report = apply(head, upload(head, again, 'RTK.xlsx').json()['id']).json()
    assert report['summary']['created'] == {'universities': 0, 'it_directions': 0, 'it_products': 0}
    assert report['summary']['updated'] == {'universities': 1, 'it_directions': 1, 'it_products': 1}
    with database(database_url) as db:
        assert db.get(University, mipt).city == 'Москва'
        assert db.get(University, mipt).short_name == 'МФТИ (новое)'
        assert db.scalar(select(ITDirection).where(ITDirection.external_id == 'DIR-001')).name == 'Анализ данных и ИИ'
        assert len(db.scalars(select(ITProduct).where(ITProduct.external_id == 'RTK-LMS-001')).all()) == 1
        assert len(db.scalars(select(VendorCompany).where(VendorCompany.name == 'ПАО «Ростелеком»')).all()) == 1


def test_a_workbook_row_without_an_id_or_name_is_an_error(head, mipt):
    book = customer_workbook(directions=[[None, 'Без кода', '', ''], ['DIR-9', None, '', '']])
    report = check(head, upload(head, book, 'RTK.xlsx').json()['id']).json()
    errors = [row for row in report['rows'] if row['status'] == 'error']
    assert [(row['sheet'], row['row_number']) for row in errors] == [('Направления', 4), ('Направления', 5)]


def test_an_ordinary_catalog_workbook_still_uses_the_column_import(head):
    book = openpyxl.Workbook()
    book.active.append(['Наименование вуза', 'Вендор', 'Программное обеспечение', 'Номер договора', 'Подписание лицензии'])
    book.active.append(['Вуз', 'РТК', 'ПО', 'Д-1', '01.03.2026'])
    buffer = io.BytesIO()
    book.save(buffer)
    body = upload(head, buffer.getvalue(), 'реестр.xlsx').json()
    assert body['kind'] == 'catalog'
    assert body['mapping']['contract_number'] == 'Номер договора'


# ---- access ----------------------------------------------------------------------------------------------------

@pytest.mark.parametrize('roles', [('crm-user',)])
def test_a_kam_cannot_upload_the_json_or_the_workbook(app, keycloak, roles):
    with TestClient(app) as client:
        login(client, keycloak, roles=roles, subject='kc-kam')
        assert upload(client, applications_json(), 'заявки.json').status_code == 403
        assert upload(client, customer_workbook(), 'RTK.xlsx').status_code == 403


def test_an_administrator_can_upload_the_json(app, keycloak):
    with TestClient(app) as client:
        login(client, keycloak, roles=('crm-admin',), subject='kc-admin')
        assert upload(client, applications_json(), 'заявки.json').status_code == 201


def test_the_archived_applications_api_tolerates_applications_without_a_learner(head, database_url):
    # The archived section stays switched off in production (D-235); its API must not break on these rows when enabled.
    apply(head, upload(head, applications_json(), 'Данные оплат.json').json()['id'])
    response = head.get('/api/v1/course-applications')
    assert response.status_code == 200, response.text
    assert {(row['external_number'], row['learner_id'], row['learner_name']) for row in response.json()} == {
        ('ORD-TEST-0001', None, ''), ('ORD-TEST-0002', None, '')}
