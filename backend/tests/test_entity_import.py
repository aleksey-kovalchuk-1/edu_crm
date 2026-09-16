"""Generic entity-import wizard (T-091/T-092): universities, university_contacts, interactions.

The contract importer's own suite (test_import_api.py, test_importer_parsing.py) is untouched and covers
`entity=contracts`; this file only exercises the three new entities and the shared job/rollback/mapping
machinery built around them.
"""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app import import_jobs, jobs  # noqa: F401 - importing import_jobs registers its JOB_HANDLERS entries
from app.models import AuditEvent, BackgroundJob, Launch, StatusChange, University, UniversityContact
from helpers import database, login

XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def workbook(headers, rows):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def head(app, keycloak):
    with TestClient(app) as client:
        login(client, keycloak, roles=('crm-supervisor',), subject='kc-head', name='Павел Демо', email='pavel.demo@demo.local')
        yield client


@pytest.fixture
def manager(app, keycloak):
    with TestClient(app) as client:
        me = login(client, keycloak, roles=('crm-user',), subject='kc-manager', name='Анна Демо', email='anna.demo@demo.local')
        client.user_id = me['user']['id']
        yield client


@pytest.fixture
def other_manager(app, keycloak):
    with TestClient(app) as client:
        me = login(client, keycloak, roles=('crm-user',), subject='kc-manager-2', name='Борис Демо', email='boris.demo@demo.local')
        client.user_id = me['user']['id']
        yield client


def upload(client, entity, content, filename='файл.xlsx'):
    return client.post('/api/v1/imports', files={'file': (filename, content, XLSX_TYPE)}, data={'entity': entity})


def uploaded(client, entity, headers, rows):
    response = upload(client, entity, workbook(headers, rows))
    assert response.status_code == 201, response.text
    return response.json()


def run_pending_jobs(database_url):
    """Drains the queue synchronously, calling the same handlers the real `worker` process would -- no
    need for a live poll loop in a test. Mirrors `app.worker.run`'s per-job dispatch."""
    engine = create_engine(database_url)
    try:
        session_factory = sessionmaker(bind=engine)
        processed = []
        while True:
            with session_factory() as db:
                job = jobs.claim_next(db)
                if job is None:
                    break
                handler = jobs.JOB_HANDLERS[job.kind]
                result = handler(db, job)
                jobs.finish(db, job, result=result)
                processed.append(job.id)
        return processed
    finally:
        engine.dispose()


def apply_and_run(client, database_url, import_id, mapping):
    response = client.post(f'/api/v1/imports/{import_id}/apply', json={'mapping': mapping})
    assert response.status_code == 202, response.text
    job_id = response.json()['job_id']
    run_pending_jobs(database_url)
    return job_id


def assign_manager(head, university_id, user_id):
    response = head.put(f'/api/v1/universities/{university_id}/managers', json={'user_ids': [user_id]})
    assert response.status_code == 200, response.text


def create_university(head, name='Волжский институт цифровых технологий', **extra):
    response = head.post('/api/v1/universities', json={'name': name, 'city': 'Волгоград', **extra})
    assert response.status_code == 201, response.text
    return response.json()


UNI_HEADERS = ['Наименование вуза', 'Город', 'Краткое название', 'Регион', 'Сайт', 'Контакт']


# ---------- universities ----------

def test_universities_import_requires_supervisor_or_admin(manager):
    response = upload(manager, 'universities', workbook(UNI_HEADERS, [['Тестовый вуз', 'Казань', '', '', '', '']]))
    assert response.status_code == 403
    assert manager.get('/api/v1/imports/fields', params={'entity': 'universities'}).status_code == 403


def test_universities_import_creates_then_updates(head, database_url):
    rows = [['Новый институт', 'Казань', 'НИ', 'Татарстан', 'https://ni.example', 'Иванов']]
    body = uploaded(head, 'universities', UNI_HEADERS, rows)
    assert body['entity'] == 'universities'
    check = head.post(f"/api/v1/imports/{body['id']}/check", json={'mapping': body['mapping']}).json()
    assert check['summary'] == {'rows': 1, 'valid': 1, 'invalid': 0, 'skipped': 0, 'with_warnings': 0, 'created': {'universities': 1}, 'updated': {'universities': 0}}

    apply_and_run(head, database_url, body['id'], body['mapping'])
    with database(database_url) as db:
        university = db.scalar(select(University).where(University.name == 'Новый институт'))
        assert (university.city, university.short_name, university.region, university.website) == ('Казань', 'НИ', 'Татарстан', 'https://ni.example')
        assert db.scalar(select(AuditEvent).where(AuditEvent.action == 'university.create', AuditEvent.entity_id == str(university.id))) is not None

    status = head.get(f"/api/v1/imports/{body['id']}").json()
    assert status['status'] == 'applied'

    second_rows = [['Новый институт', 'Иннополис', 'НИ', 'Татарстан', 'https://ni.example', 'Иванов']]
    second = uploaded(head, 'universities', UNI_HEADERS, second_rows)
    check2 = head.post(f"/api/v1/imports/{second['id']}/check", json={'mapping': second['mapping']}).json()
    assert check2['rows'][0]['action'] == 'update'
    assert check2['summary']['created'] == {'universities': 0}
    apply_and_run(head, database_url, second['id'], second['mapping'])
    with database(database_url) as db:
        assert db.scalar(select(University.city).where(University.name == 'Новый институт')) == 'Иннополис'
        assert db.scalar(select(University).where(University.name == 'Новый институт')) is not None
        # Exactly one row: re-import updates in place, does not duplicate.
        assert len(db.scalars(select(University).where(University.name == 'Новый институт')).all()) == 1


def test_universities_import_missing_required_field(head):
    body = uploaded(head, 'universities', UNI_HEADERS, [['', 'Казань', '', '', '', '']])
    report = head.post(f"/api/v1/imports/{body['id']}/check", json={'mapping': body['mapping']}).json()
    assert report['rows'][0]['status'] == 'error'
    assert any('Наименование вуза' in message for message in report['rows'][0]['errors'])


def test_universities_import_invalid_website(head):
    body = uploaded(head, 'universities', UNI_HEADERS, [['Институт Х', 'Казань', '', '', 'not-a-url', '']])
    report = head.post(f"/api/v1/imports/{body['id']}/check", json={'mapping': body['mapping']}).json()
    assert report['rows'][0]['status'] == 'error'
    assert any('http' in message for message in report['rows'][0]['errors'])


def test_universities_import_duplicate_row_in_file(head, database_url):
    rows = [
        ['Дубль Вуз', 'Казань', '', '', '', ''],
        ['Дубль Вуз', 'Уфа', '', '', '', ''],
    ]
    body = uploaded(head, 'universities', UNI_HEADERS, rows)
    report = head.post(f"/api/v1/imports/{body['id']}/check", json={'mapping': body['mapping']}).json()
    assert [row['status'] for row in report['rows']] == ['skipped', 'ok']
    apply_and_run(head, database_url, body['id'], body['mapping'])
    with database(database_url) as db:
        assert db.scalar(select(University.city).where(University.name == 'Дубль Вуз')) == 'Уфа'
        assert len(db.scalars(select(University).where(University.name == 'Дубль Вуз')).all()) == 1


# ---------- university_contacts ----------

CONTACT_HEADERS = ['Наименование вуза', 'ФИО', 'Должность', 'Email', 'Телефон', 'Комментарий']


def test_contacts_import_creates_and_respects_scope(head, manager, database_url):
    uni_in_scope = create_university(head, 'Вуз В Скоупе')
    uni_out_of_scope = create_university(head, 'Вуз Вне Скоупа')
    assign_manager(head, uni_in_scope['id'], manager.user_id)

    rows = [
        [uni_in_scope['name'], 'Иванов Иван', 'Декан', 'ivanov@example.com', '', ''],
        [uni_out_of_scope['name'], 'Петров Пётр', 'Декан', '', '', ''],
    ]
    body = uploaded(manager, 'university_contacts', CONTACT_HEADERS, rows)
    report = manager.post(f"/api/v1/imports/{body['id']}/check", json={'mapping': body['mapping']}).json()
    assert report['rows'][0]['status'] == 'ok'
    assert report['rows'][1]['status'] == 'error'
    assert any('не назначено' in message for message in report['rows'][1]['errors'])

    apply_and_run(manager, database_url, body['id'], body['mapping'])
    with database(database_url) as db:
        assert db.scalar(select(UniversityContact).where(UniversityContact.full_name == 'Иванов Иван')) is not None
        assert db.scalar(select(UniversityContact).where(UniversityContact.full_name == 'Петров Пётр')) is None


def test_contacts_import_invalid_email(head):
    uni = create_university(head, 'Вуз Почты')
    body = uploaded(head, 'university_contacts', CONTACT_HEADERS, [[uni['name'], 'Сидоров Сидор', '', 'bad@', '', '']])
    report = head.post(f"/api/v1/imports/{body['id']}/check", json={'mapping': body['mapping']}).json()
    assert report['rows'][0]['status'] == 'error'
    assert any('почты' in message for message in report['rows'][0]['errors'])


def test_contacts_import_university_not_found(head):
    body = uploaded(head, 'university_contacts', CONTACT_HEADERS, [['Несуществующий вуз', 'Кто-то', '', '', '', '']])
    report = head.post(f"/api/v1/imports/{body['id']}/check", json={'mapping': body['mapping']}).json()
    assert report['rows'][0]['status'] == 'error'
    assert any('не найдено' in message for message in report['rows'][0]['errors'])


# ---------- interactions ----------

LAUNCH_HEADERS = ['Наименование вуза', 'Программа', 'Продукт', 'Ответственный', 'Число студентов', 'Срок']


def test_interactions_import_creates_launch_with_default_workflow(head, database_url):
    uni = create_university(head, 'Вуз Взаимодействий')
    rows = [[uni['name'], 'Python для начинающих', 'Учебная среда', 'Иванов И.И.', 25, '01.12.2026']]
    body = uploaded(head, 'interactions', LAUNCH_HEADERS, rows)
    apply_and_run(head, database_url, body['id'], body['mapping'])

    with database(database_url) as db:
        launch = db.scalar(select(Launch).where(Launch.program == 'Python для начинающих'))
        assert launch is not None
        assert (launch.product, launch.owner, launch.students) == ('Учебная среда', 'Иванов И.И.', 25)
        assert launch.stage == 0
        assert db.scalar(select(StatusChange).where(StatusChange.launch_id == launch.id)) is not None
        assert db.scalar(select(AuditEvent).where(AuditEvent.action == 'launch.create', AuditEvent.entity_id == str(launch.id))) is not None


def test_interactions_import_duplicate_create_only_no_update(head, database_url):
    uni = create_university(head, 'Вуз Повтора')
    rows = [
        [uni['name'], 'Курс А', 'Продукт', 'Менеджер', 10, '01.12.2026'],
        [uni['name'], 'Курс А', 'Продукт', 'Менеджер', 10, '01.12.2026'],
    ]
    body = uploaded(head, 'interactions', LAUNCH_HEADERS, rows)
    report = head.post(f"/api/v1/imports/{body['id']}/check", json={'mapping': body['mapping']}).json()
    assert [row['status'] for row in report['rows']] == ['skipped', 'ok']
    apply_and_run(head, database_url, body['id'], body['mapping'])
    with database(database_url) as db:
        assert len(db.scalars(select(Launch).where(Launch.program == 'Курс А')).all()) == 1


def test_interactions_import_out_of_scope_university_rejected(manager, head, database_url):
    uni = create_university(head, 'Вуз Не Назначен Менеджеру')
    body = uploaded(manager, 'interactions', LAUNCH_HEADERS, [[uni['name'], 'Курс', 'Продукт', 'Менеджер', 5, '01.12.2026']])
    report = manager.post(f"/api/v1/imports/{body['id']}/check", json={'mapping': body['mapping']}).json()
    assert report['rows'][0]['status'] == 'error'
    apply_and_run(manager, database_url, body['id'], body['mapping'])
    with database(database_url) as db:
        assert db.scalar(select(Launch).where(Launch.program == 'Курс')) is None


def test_interactions_import_invalid_deadline_and_students(head):
    uni = create_university(head, 'Вуз Валидации')
    rows = [[uni['name'], 'Курс', 'Продукт', 'Менеджер', 'много', 'не дата']]
    body = uploaded(head, 'interactions', LAUNCH_HEADERS, rows)
    report = head.post(f"/api/v1/imports/{body['id']}/check", json={'mapping': body['mapping']}).json()
    assert report['rows'][0]['status'] == 'error'
    assert len(report['rows'][0]['errors']) == 2


# ---------- apply idempotency ----------

def test_reapply_is_idempotent_no_duplicate_job_or_write(head, database_url):
    uni = create_university(head, 'Вуз Идемпотентности')
    body = uploaded(head, 'interactions', LAUNCH_HEADERS, [[uni['name'], 'Idem', 'Продукт', 'Менеджер', 1, '01.12.2026']])

    first = head.post(f"/api/v1/imports/{body['id']}/apply", json={'mapping': body['mapping']})
    assert first.status_code == 202
    # A second apply call before the job has run must reuse the same job, not enqueue another one.
    second = head.post(f"/api/v1/imports/{body['id']}/apply", json={'mapping': body['mapping']})
    assert second.status_code == 202
    assert second.json()['job_id'] == first.json()['job_id']
    with database(database_url) as db:
        assert db.scalar(select(BackgroundJob).where(BackgroundJob.kind == 'import_apply')) is not None
        count = len(db.scalars(select(BackgroundJob).where(BackgroundJob.kind == 'import_apply')).all())
        assert count == 1

    run_pending_jobs(database_url)
    third = head.post(f"/api/v1/imports/{body['id']}/apply", json={'mapping': body['mapping']})
    assert third.status_code == 409
    with database(database_url) as db:
        assert len(db.scalars(select(Launch).where(Launch.program == 'Idem')).all()) == 1


# ---------- job status endpoint ----------

def test_job_status_endpoint_scoping(head, manager, other_manager, database_url):
    uni = create_university(head, 'Вуз Задач')
    assign_manager(head, uni['id'], manager.user_id)
    body = uploaded(manager, 'interactions', LAUNCH_HEADERS, [[uni['name'], 'Задача', 'Продукт', 'Менеджер', 1, '01.12.2026']])
    job_id = manager.post(f"/api/v1/imports/{body['id']}/apply", json={'mapping': body['mapping']}).json()['job_id']

    own = manager.get(f'/api/v1/jobs/{job_id}')
    assert own.status_code == 200
    assert own.json()['kind'] == 'import_apply'

    assert other_manager.get(f'/api/v1/jobs/{job_id}').status_code == 404
    assert head.get(f'/api/v1/jobs/{job_id}').status_code == 200

    run_pending_jobs(database_url)
    finished = manager.get(f'/api/v1/jobs/{job_id}').json()
    assert finished['status'] == 'succeeded'
    assert finished['result']['summary']['created'] == {'interactions': 1}


# ---------- saved mapping profiles ----------

def test_saved_mapping_profile_create_reuse_and_duplicate_conflict(head):
    body = uploaded(head, 'universities', UNI_HEADERS, [['Профильный вуз', 'Казань', '', '', '', '']])
    create = head.post('/api/v1/import-mappings', json={'entity': 'universities', 'name': 'Стандартный профиль', 'mapping': body['mapping']})
    assert create.status_code == 201, create.text
    saved = create.json()
    assert saved['entity'] == 'universities' and saved['mapping'] == body['mapping']

    listed = head.get('/api/v1/import-mappings', params={'entity': 'universities'})
    assert listed.status_code == 200
    assert any(item['name'] == 'Стандартный профиль' for item in listed.json())

    duplicate = head.post('/api/v1/import-mappings', json={'entity': 'universities', 'name': 'Стандартный профиль', 'mapping': body['mapping']})
    assert duplicate.status_code == 409

    # Applying the saved mapping to a fresh upload works exactly like a freshly-suggested one.
    second = uploaded(head, 'universities', UNI_HEADERS, [['Профильный вуз 2', 'Уфа', '', '', '', '']])
    check = head.post(f"/api/v1/imports/{second['id']}/check", json={'mapping': saved['mapping']}).json()
    assert check['rows'][0]['status'] == 'ok'


# ---------- error report download ----------

def test_error_report_download_is_409_before_a_report_exists(head, database_url):
    rows = [['', 'Казань', '', '', '', '']]
    body = uploaded(head, 'universities', UNI_HEADERS, rows)
    # `check` is a dry run and rolls back; it does not persist a report, so nothing is downloadable yet.
    head.post(f"/api/v1/imports/{body['id']}/check", json={'mapping': body['mapping']})
    assert head.get(f"/api/v1/imports/{body['id']}/errors.xlsx").status_code == 409

    # Apply enqueues the job but the report is only written once the worker has run it.
    apply_response = head.post(f"/api/v1/imports/{body['id']}/apply", json={'mapping': body['mapping']})
    assert apply_response.status_code == 202
    assert head.get(f"/api/v1/imports/{body['id']}/errors.xlsx").status_code == 409
    run_pending_jobs(database_url)
    assert head.get(f"/api/v1/imports/{body['id']}/errors.xlsx").status_code == 200


def test_error_report_download_after_apply(head, database_url):
    rows = [['Отчётный вуз 2', 'Казань', '', '', '', ''], ['', 'Казань', '', '', '', '']]
    body = uploaded(head, 'universities', UNI_HEADERS, rows)
    apply_and_run(head, database_url, body['id'], body['mapping'])

    response = head.get(f"/api/v1/imports/{body['id']}/errors.xlsx")
    assert response.status_code == 200
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats')
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = book.active
    values = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    assert values[0] == ('Строка', 'Тип', 'Сообщение')
    assert any(row[1] == 'Ошибка' for row in values[1:])


# ---------- rollback ----------

def test_rollback_undoes_untouched_created_row(head, database_url):
    body = uploaded(head, 'universities', UNI_HEADERS, [['Откатной вуз', 'Казань', '', '', '', '']])
    apply_and_run(head, database_url, body['id'], body['mapping'])
    with database(database_url) as db:
        assert db.scalar(select(University).where(University.name == 'Откатной вуз')) is not None

    rollback = head.post(f"/api/v1/imports/{body['id']}/rollback")
    assert rollback.status_code == 202, rollback.text
    run_pending_jobs(database_url)

    with database(database_url) as db:
        assert db.scalar(select(University).where(University.name == 'Откатной вуз')) is None

    job_id = rollback.json()['job_id']
    result = head.get(f'/api/v1/jobs/{job_id}').json()
    assert result['result']['summary'] == {'total': 1, 'rolled_back': 1, 'not_rolled_back': 0}


def test_rollback_reports_since_modified_row_as_not_rolled_back(head, database_url):
    body = uploaded(head, 'universities', UNI_HEADERS, [['Изменённый после импорта вуз', 'Казань', '', '', '', '']])
    apply_and_run(head, database_url, body['id'], body['mapping'])
    with database(database_url) as db:
        university_id = db.scalar(select(University.id).where(University.name == 'Изменённый после импорта вуз'))

    changed = head.patch(f'/api/v1/universities/{university_id}', json={'city': 'Другой город'})
    assert changed.status_code == 200

    rollback = head.post(f"/api/v1/imports/{body['id']}/rollback")
    assert rollback.status_code == 202
    run_pending_jobs(database_url)

    with database(database_url) as db:
        university = db.get(University, university_id)
        assert university is not None
        assert university.city == 'Другой город'  # not reverted

    job_id = rollback.json()['job_id']
    result = head.get(f'/api/v1/jobs/{job_id}').json()['result']
    assert result['summary'] == {'total': 1, 'rolled_back': 0, 'not_rolled_back': 1}
    assert result['items'][0]['reason']


CONTRACT_HEADERS = ['Наименование вуза', 'Вендор', 'Программное обеспечение', 'Номер договора', 'Подписание лицензии']
CONTRACT_ROWS = [['Договорной вуз', 'РТК ИТ', 'Учебная среда', 'Д-900', '15.01.2026']]


def test_rollback_not_supported_for_contract_imports(head, database_url):
    # entity defaults to 'contracts' when the form omits it, exactly like the pre-existing endpoint.
    upload_response = head.post('/api/v1/imports', files={'file': ('реестр.xlsx', workbook(CONTRACT_HEADERS, CONTRACT_ROWS), XLSX_TYPE)})
    assert upload_response.status_code == 201, upload_response.text
    body = upload_response.json()
    assert body['entity'] == 'contracts'
    applied = head.post(f"/api/v1/imports/{body['id']}/apply", json={'mapping': body['mapping']})
    assert applied.status_code == 200  # contracts still apply synchronously, unchanged

    rollback = head.post(f"/api/v1/imports/{body['id']}/rollback")
    assert rollback.status_code == 409
