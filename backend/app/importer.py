"""Reading and interpreting catalog upload files (docs/design/import.md).

Pure functions without database access: the file is validated and parsed into rows, headers are matched to
CRM fields, and each row is checked and converted into typed values.
"""
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl
import xlrd
from openpyxl.utils.datetime import from_excel

from .catalog_routes import TRANSFER_STATUS_LABELS, one_year_after

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_DATA_ROWS = 5000
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_COLUMNS = 50
HEADER_SEARCH_ROWS = 10
XLSX_SIGNATURE = b'PK\x03\x04'
XLS_SIGNATURE = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'

# CRM field -> (required, maximum length, header synonyms already normalised by normalize_header).
FIELDS = {
    'university_name': (True, 200, ['наименование вуза', 'вуз', 'учебное заведение', 'наименование учебного заведения', 'university name', 'university']),
    'vendor': (True, 200, ['вендор', 'производитель', 'vendor']),
    'software': (True, 200, ['по', 'программное обеспечение', 'продукт', 'ит продукт', 'software', 'it product']),
    'contract_number': (True, 100, ['номер договора', 'договор', 'contract number', 'contract']),
    'license_signed_at': (True, None, ['подписание лицензии', 'дата подписания', 'дата подписания лицензии', 'signing the license', 'signed at']),
    'license_valid_until': (False, None, ['срок действия лицензии', 'действует до', 'срок действия', 'license validity period', 'valid until']),
    'transfer_status': (False, None, ['статус передачи', 'transfer status']),
    'manager_full_name': (False, 200, ['фио менеджера', 'менеджер', 'manager', 'manager full name', 'manager s full name']),
    'university_contacts': (False, None, ['ответственные от вуза', 'ответственные лица вуза', 'ответственные лица от вуза', 'responsible persons from the university', 'responsible persons']),
    'comment': (False, 2000, ['комментарий', 'примечание', 'comment']),
    'it_directions': (False, None, ['ит направления', 'ит направление', 'направление', 'it directions', 'it direction']),
}
REQUIRED_FIELDS = [name for name, (required, _, _) in FIELDS.items() if required]
FIELD_LABELS = {
    'university_name': 'Наименование вуза',
    'vendor': 'Вендор',
    'software': 'Программное обеспечение',
    'contract_number': 'Номер договора',
    'license_signed_at': 'Подписание лицензии',
    'license_valid_until': 'Срок действия лицензии',
    'transfer_status': 'Статус передачи',
    'manager_full_name': 'ФИО менеджера',
    'university_contacts': 'Ответственные от вуза',
    'comment': 'Комментарий',
    'it_directions': 'ИТ-направления',
}
LIST_ITEM_MAX_LENGTH = {'university_contacts': 200, 'it_directions': 120}
DATE_FORMATS = ('%d.%m.%Y', '%Y-%m-%d', '%d/%m/%Y', '%d.%m.%y')


class ImportFileError(Exception):
    """The upload cannot be used; the message is shown to the user."""


class UnsupportedFileType(ImportFileError):
    """The upload is not an Excel workbook."""


@dataclass
class ParsedSheet:
    header_row: int
    headers: list[str]
    rows: list[tuple[int, list]]


@dataclass
class RowResult:
    row_number: int
    values: dict = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def is_valid(self):
        return not self.errors


def normalize_header(value):
    text = str(value or '').strip().lower().replace('ё', 'е').replace('№', ' номер ')
    text = re.sub(r'[^\w\s]', ' ', text)
    return re.sub(r'\s+', ' ', text).strip()


def normalize_name(value):
    """Collapses whitespace; used to match names typed slightly differently in spreadsheets."""
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def name_key(value):
    return normalize_name(value).casefold().replace('ё', 'е')


def read_upload(filename, content):
    name = (filename or '').lower()
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError('Файл больше 10 МБ')
    if name.endswith('.xlsx') and content.startswith(XLSX_SIGNATURE):
        cells = _read_xlsx(content)
    elif name.endswith('.xls') and content.startswith(XLS_SIGNATURE):
        cells = _read_xls(content)
    else:
        raise UnsupportedFileType('Поддерживаются только файлы Excel .xls и .xlsx')
    return _split_header(cells)


def _read_xlsx(content):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            # A small archive can expand to gigabytes; check before any XML is parsed.
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ImportFileError('Файл слишком большой после распаковки')
    except zipfile.BadZipFile as error:
        raise ImportFileError('Файл .xlsx повреждён') from error
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises many unrelated types for malformed files
        raise ImportFileError('Не удалось прочитать файл .xlsx') from error
    try:
        if not workbook.worksheets:
            raise ImportFileError('В файле нет листов')
        rows = []
        for row in workbook.worksheets[0].iter_rows(values_only=True):
            rows.append(list(row[:MAX_COLUMNS]))
            _check_row_count(len(rows))
        return rows
    finally:
        workbook.close()


def _read_xls(content):
    try:
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception as error:  # xlrd raises several error types for malformed files
        raise ImportFileError('Не удалось прочитать файл .xls') from error
    try:
        if book.nsheets == 0:
            raise ImportFileError('В файле нет листов')
        sheet = book.sheet_by_index(0)
        _check_row_count(sheet.nrows)
        rows = []
        for row_index in range(sheet.nrows):
            row = []
            for column_index in range(min(sheet.ncols, MAX_COLUMNS)):
                cell = sheet.cell(row_index, column_index)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                    row.append(int(cell.value))
                else:
                    row.append(cell.value)
            rows.append(row)
        return rows
    finally:
        book.release_resources()


def _check_row_count(count):
    if count > MAX_DATA_ROWS + HEADER_SEARCH_ROWS + 1:
        raise ImportFileError(f'В файле больше {MAX_DATA_ROWS} строк данных')


def _is_empty(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _split_header(cells):
    known = {synonym for _, _, synonyms in FIELDS.values() for synonym in synonyms}
    header_index = None
    for index, row in enumerate(cells[:HEADER_SEARCH_ROWS]):
        # A title row may precede the table; the header is the first row naming at least two known fields.
        if sum(1 for value in row if normalize_header(value) in known) >= 2:
            header_index = index
            break
    if header_index is None:
        header_index = next((index for index, row in enumerate(cells) if any(not _is_empty(value) for value in row)), None)
    if header_index is None:
        raise ImportFileError('Файл пустой')

    raw_headers = cells[header_index]
    headers = [normalize_name(value) if not _is_empty(value) else f'Столбец {position + 1}' for position, value in enumerate(raw_headers)]
    rows = [
        (header_index + offset + 2, (row + [None] * len(headers))[:len(headers)])
        for offset, row in enumerate(cells[header_index + 1:])
        if any(not _is_empty(value) for value in row)
    ]
    if not rows:
        raise ImportFileError('В файле нет строк с данными')
    if len(rows) > MAX_DATA_ROWS:
        raise ImportFileError(f'В файле больше {MAX_DATA_ROWS} строк данных')
    return ParsedSheet(header_row=header_index + 1, headers=headers, rows=rows)


def suggest_mapping(headers):
    mapping, used = {}, set()
    for field_name, (_, _, synonyms) in FIELDS.items():
        match = next((header for header in headers if header not in used and normalize_header(header) in synonyms), None)
        mapping[field_name] = match
        if match is not None:
            used.add(match)
    return mapping


def validate_mapping(headers, mapping):
    """Returns a list of problems; an empty list means the mapping can be used."""
    problems = []
    unknown_fields = sorted(set(mapping) - set(FIELDS))
    if unknown_fields:
        problems.append(f'Неизвестные поля: {", ".join(unknown_fields)}')
    for field_name in REQUIRED_FIELDS:
        if not mapping.get(field_name):
            problems.append(f'Не выбран столбец для поля «{FIELD_LABELS[field_name]}»')
    chosen = [header for header in mapping.values() if header]
    missing_headers = sorted({header for header in chosen if header not in headers})
    if missing_headers:
        problems.append(f'В файле нет столбцов: {", ".join(missing_headers)}')
    duplicates = sorted({header for header in chosen if chosen.count(header) > 1})
    if duplicates:
        problems.append(f'Один столбец выбран для нескольких полей: {", ".join(duplicates)}')
    return problems


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # A date typed into a cell without date formatting arrives as an Excel serial number.
        if 1 <= value <= 2958465:
            return from_excel(value).date()
        return None
    text = normalize_name(value)
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    return None


STATUS_BY_TEXT = {
    **{code: code for code in TRANSFER_STATUS_LABELS},
    **{label.casefold().replace('ё', 'е'): code for code, label in TRANSFER_STATUS_LABELS.items()},
    'в процессе': 'in_progress',
    'передана': 'transferred',
    'отменена': 'cancelled',
}


def split_list(value):
    return [normalize_name(item) for item in re.split(r'[;,\n]', str(value or '')) if normalize_name(item)]


def _text(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return normalize_name(value) if not _is_empty(value) else ''


def interpret_row(row_number, record):
    """`record` maps CRM field names to raw cell values; returns typed values plus errors and warnings."""
    result = RowResult(row_number=row_number)
    values = result.values

    for field_name, (required, max_length, _) in FIELDS.items():
        if field_name in ('license_signed_at', 'license_valid_until', 'university_contacts', 'it_directions'):
            continue
        text = _text(record.get(field_name))
        if required and not text:
            result.errors.append(f'Не заполнено поле «{FIELD_LABELS[field_name]}»')
        if max_length and len(text) > max_length:
            result.errors.append(f'Поле «{FIELD_LABELS[field_name]}» длиннее {max_length} символов')
        values[field_name] = text

    raw_signed = record.get('license_signed_at')
    signed_at = parse_date(raw_signed)
    if _is_empty(raw_signed):
        result.errors.append(f'Не заполнено поле «{FIELD_LABELS["license_signed_at"]}»')
    elif signed_at is None:
        result.errors.append(f'Не удалось распознать дату подписания: «{_text(raw_signed)}»')
    values['license_signed_at'] = signed_at

    raw_valid = record.get('license_valid_until')
    if _is_empty(raw_valid):
        values['license_valid_until'] = one_year_after(signed_at) if signed_at else None
    else:
        valid_until = parse_date(raw_valid)
        if valid_until is None:
            result.errors.append(f'Не удалось распознать срок действия лицензии: «{_text(raw_valid)}»')
        elif signed_at and valid_until < signed_at:
            result.errors.append('Срок действия лицензии раньше даты подписания')
        values['license_valid_until'] = valid_until

    status_text = values.pop('transfer_status')
    if status_text:
        status = STATUS_BY_TEXT.get(status_text.casefold().replace('ё', 'е'))
        if status is None:
            result.errors.append(f'Неизвестный статус передачи: «{status_text}»')
        values['transfer_status'] = status
    else:
        values['transfer_status'] = 'not_started'

    for field_name in ('university_contacts', 'it_directions'):
        items = split_list(record.get(field_name))
        too_long = [item for item in items if len(item) > LIST_ITEM_MAX_LENGTH[field_name]]
        if too_long:
            result.errors.append(f'В поле «{FIELD_LABELS[field_name]}» значение длиннее {LIST_ITEM_MAX_LENGTH[field_name]} символов')
        values[field_name] = list(dict.fromkeys(items))
    return result


def records(sheet, mapping):
    """Yields (row number, {field: raw value}) using the chosen header for each field."""
    positions = {field_name: sheet.headers.index(header) for field_name, header in mapping.items() if header}
    for row_number, cells in sheet.rows:
        yield row_number, {field_name: cells[position] for field_name, position in positions.items()}
